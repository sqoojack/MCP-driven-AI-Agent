from pptx.util import Pt
from pptx.dml.color import RGBColor

class TextRunFactory:
    def __init__(self, paragraph):
        self.paragraph = paragraph

    def add_icon(self, icon_text, font_size=30, color = RGBColor(0, 102, 204)):
        run = self.paragraph.add_run()
        run.font.name = "Segoe UI Emoji"
        run.text = icon_text + "\n"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        return run

    def add_id(self, node_id, size=13):
        run = self.paragraph.add_run()
        run.text = node_id + "\n"
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run

    def add_add(self, text, size=10):
        run = self.paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run
